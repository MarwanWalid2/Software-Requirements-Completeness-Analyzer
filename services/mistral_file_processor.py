import os
import logging
import tempfile
import json
from pathlib import Path
from typing import Tuple, Optional
from mistralai import Mistral, DocumentURLChunk
from mistralai.models import OCRResponse
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.enums import TA_LEFT
import traceback
from config import get_mistral_api_key

logger = logging.getLogger(__name__)

class MistralFileProcessor:
    """
    File processor that uses Mistral AI's OCR service to extract text from documents.
    Supports various file formats by converting them to PDF first if needed.
    """
    
    def __init__(self):
        """Initialize the Mistral AI client"""
        api_key = get_mistral_api_key()
        if api_key and api_key != "dummy_key_replace_me":
            self.client = Mistral(api_key=api_key)
            self.enabled = True
            logger.info("Mistral AI client initialized successfully")
        else:
            self.client = None
            self.enabled = False
            logger.warning("Mistral AI client not available - no valid API key")
    
    def process_file(self, file_path: str, filename: str) -> Tuple[str, dict]:
        """
        Process a file using Mistral AI's OCR service.
        
        Args:
            file_path (str): Path to the uploaded file
            filename (str): Original filename
            
        Returns:
            Tuple[str, dict]: (extracted_text, metadata)
        """
        if not self.enabled:
            raise ValueError("Mistral AI OCR service is not available. Please check your MISTRAL_API_KEY.")
        
        try:
            logger.info(f"Processing file: {filename}")
            
            # Get file extension
            file_extension = Path(filename).suffix.lower()
            
            # Convert to PDF if necessary
            pdf_path = self._convert_to_pdf_if_needed(file_path, filename, file_extension)
            
            # Upload PDF to Mistral's OCR service
            uploaded_file = self._upload_file_to_mistral(pdf_path, filename)
            
            # Get signed URL for the uploaded file
            signed_url = self.client.files.get_signed_url(
                file_id=uploaded_file.id, 
                expiry=1
            )
            
            # Process PDF with OCR
            logger.info("Processing file with Mistral AI OCR...")
            ocr_response = self.client.ocr.process(
                document=DocumentURLChunk(document_url=signed_url.url),
                model="mistral-ocr-latest",
                include_image_base64=True
            )
            
            # Extract text from OCR response
            extracted_text = self._extract_text_from_ocr_response(ocr_response)
            
            # Create metadata
            metadata = {
                "original_filename": filename,
                "file_extension": file_extension,
                "processed_with": "mistral_ocr",
                "pages_processed": len(ocr_response.pages) if ocr_response.pages else 0,
                "has_images": any(page.images for page in ocr_response.pages) if ocr_response.pages else False
            }
            
            # Clean up temporary files
            if pdf_path != file_path and os.path.exists(pdf_path):
                os.remove(pdf_path)
                logger.debug(f"Cleaned up temporary PDF: {pdf_path}")
            
            logger.info(f"Successfully processed {filename} - extracted {len(extracted_text)} characters")
            return extracted_text, metadata
            
        except Exception as e:
            logger.error(f"Error processing file {filename}: {str(e)}")
            logger.error(traceback.format_exc())
            raise
    
    def _convert_to_pdf_if_needed(self, file_path: str, filename: str, file_extension: str) -> str:
        """
        Convert file to PDF if it's not already a PDF.
        
        Args:
            file_path (str): Path to the original file
            filename (str): Original filename
            file_extension (str): File extension
            
        Returns:
            str: Path to PDF file (original or converted)
        """
        if file_extension == '.pdf':
            logger.debug("File is already PDF, no conversion needed")
            return file_path
        
        logger.info(f"Converting {file_extension} file to PDF")
        
        # Create temporary PDF file
        temp_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
        temp_pdf_path = temp_pdf.name
        temp_pdf.close()
        
        try:
            if file_extension in ['.txt', '.md']:
                self._convert_text_to_pdf(file_path, temp_pdf_path)
            elif file_extension in ['.docx']:
                self._convert_docx_to_pdf(file_path, temp_pdf_path)
            elif file_extension in ['.doc']:
                # For .doc files, try to read as text first
                self._convert_legacy_doc_to_pdf(file_path, temp_pdf_path)
            else:
                raise ValueError(f"Unsupported file format: {file_extension}")
            
            logger.debug(f"Successfully converted {filename} to PDF: {temp_pdf_path}")
            return temp_pdf_path
            
        except Exception as e:
            # Clean up on error
            if os.path.exists(temp_pdf_path):
                os.remove(temp_pdf_path)
            raise ValueError(f"Failed to convert {filename} to PDF: {str(e)}")
    
    def _convert_text_to_pdf(self, text_file_path: str, pdf_path: str):
        """Convert text file to PDF"""
        try:
            # Read text content
            with open(text_file_path, 'r', encoding='utf-8', errors='replace') as f:
                content = f.read()
            
            # Create PDF
            doc = SimpleDocTemplate(pdf_path, pagesize=letter, 
                                  leftMargin=0.75*inch, rightMargin=0.75*inch,
                                  topMargin=1*inch, bottomMargin=1*inch)
            
            # Define styles
            styles = getSampleStyleSheet()
            normal_style = ParagraphStyle(
                'CustomNormal',
                parent=styles['Normal'],
                fontSize=10,
                spaceAfter=12,
                alignment=TA_LEFT
            )
            
            # Split content into paragraphs
            paragraphs = content.split('\n\n')
            story = []
            
            for para_text in paragraphs:
                if para_text.strip():
                    # Replace newlines within paragraphs with spaces
                    para_text = para_text.replace('\n', ' ').strip()
                    para = Paragraph(para_text, normal_style)
                    story.append(para)
                    story.append(Spacer(1, 12))
            
            # Build PDF
            doc.build(story)
            logger.debug("Successfully converted text to PDF")
            
        except Exception as e:
            raise ValueError(f"Error converting text to PDF: {str(e)}")
    
    def _convert_docx_to_pdf(self, docx_path: str, pdf_path: str):
        """Convert DOCX file to PDF"""
        try:
            import docx
            
            # Read DOCX content
            doc = docx.Document(docx_path)
            content = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    content.append(paragraph.text)
            
            # Join all paragraphs
            full_text = '\n\n'.join(content)
            
            # Create temporary text file and convert to PDF
            temp_txt = tempfile.NamedTemporaryFile(mode='w', suffix='.txt', 
                                                 encoding='utf-8', delete=False)
            temp_txt.write(full_text)
            temp_txt.close()
            
            try:
                self._convert_text_to_pdf(temp_txt.name, pdf_path)
            finally:
                os.remove(temp_txt.name)
            
            logger.debug("Successfully converted DOCX to PDF")
            
        except ImportError:
            raise ValueError("python-docx library not installed. Cannot process DOCX files.")
        except Exception as e:
            raise ValueError(f"Error converting DOCX to PDF: {str(e)}")
    
    def _convert_legacy_doc_to_pdf(self, doc_path: str, pdf_path: str):
        """Convert legacy DOC file to PDF by reading as binary and extracting text"""
        try:
            # Try to extract text from .doc file
            # This is a basic approach - for better results, consider using python-docx2txt or antiword
            with open(doc_path, 'rb') as f:
                raw_content = f.read()
            
            # Try to decode as text (this won't work perfectly for .doc files)
            try:
                # Simple text extraction - may not work well for complex .doc files
                content = raw_content.decode('utf-8', errors='ignore')
                # Filter out non-printable characters
                content = ''.join(char for char in content if char.isprintable() or char.isspace())
            except:
                content = "Could not extract text from legacy .doc file. Please convert to .docx or .pdf format."
            
            # Create temporary text file and convert to PDF
            temp_txt = tempfile.NamedTemporaryFile(mode='w', suffix='.txt', 
                                                 encoding='utf-8', delete=False)
            temp_txt.write(content)
            temp_txt.close()
            
            try:
                self._convert_text_to_pdf(temp_txt.name, pdf_path)
            finally:
                os.remove(temp_txt.name)
            
            logger.warning("Converted legacy .doc file using basic text extraction. For better results, please use .docx format.")
            
        except Exception as e:
            raise ValueError(f"Error converting legacy DOC to PDF: {str(e)}")
    
    def _upload_file_to_mistral(self, file_path: str, original_filename: str):
        """Upload file to Mistral AI OCR service"""
        try:
            with open(file_path, 'rb') as f:
                file_content = f.read()
            
            uploaded_file = self.client.files.upload(
                file={
                    "file_name": Path(original_filename).stem,
                    "content": file_content,
                },
                purpose="ocr",
            )
            
            logger.debug(f"Successfully uploaded file to Mistral AI: {uploaded_file.id}")
            return uploaded_file
            
        except Exception as e:
            raise ValueError(f"Error uploading file to Mistral AI: {str(e)}")
    
    def _extract_text_from_ocr_response(self, ocr_response: OCRResponse) -> str:
        """
        Extract combined text from OCR response.
        
        Args:
            ocr_response: Response from Mistral AI OCR processing
            
        Returns:
            str: Combined text from all pages
        """
        try:
            combined_text = self._get_combined_markdown(ocr_response)
            
            # Clean up the markdown for plain text extraction if needed
            # For now, we'll return the markdown as-is since it contains the text content
            return combined_text
            
        except Exception as e:
            logger.error(f"Error extracting text from OCR response: {str(e)}")
            return "Error: Could not extract text from OCR response"
    
    def _replace_images_in_markdown(self, markdown_str: str, images_dict: dict) -> str:
        """
        Replace image placeholders in markdown with base64-encoded images.
        
        Args:
            markdown_str: Markdown text containing image placeholders
            images_dict: Dictionary mapping image IDs to base64 strings
            
        Returns:
            Markdown text with images replaced by base64 data
        """
        for img_name, base64_str in images_dict.items():
            markdown_str = markdown_str.replace(
                f"![{img_name}]({img_name})",
                f"![{img_name}](data:image/png;base64,{base64_str})"
            )
        return markdown_str
    
    def _get_combined_markdown(self, ocr_response: OCRResponse) -> str:
        """
        Combine OCR text and images into a single markdown document.
        
        Args:
            ocr_response: Response from OCR processing containing text and images
            
        Returns:
            Combined markdown string with embedded images
        """
        markdowns = []
        
        # Extract images from each page
        for page in ocr_response.pages:
            image_data = {}
            for img in page.images:
                image_data[img.id] = img.image_base64
            
            # Replace image placeholders with actual images
            page_markdown = self._replace_images_in_markdown(page.markdown, image_data)
            markdowns.append(page_markdown)
        
        return "\n\n".join(markdowns)

# Global instance
mistral_processor = MistralFileProcessor()